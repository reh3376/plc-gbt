"""
DOCX Text Extraction for Control Theory Documents
Extracts text content from Word documents for intelligent processing
"""

import sys
from pathlib import Path

try:
    from docx import Document
except ImportError:
    print("python-docx not found. Installing...")
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    from docx import Document

def extract_docx_text(docx_path: str) -> str:
    """
    Extract all text from a DOCX file including paragraphs and tables
    
    Args:
        docx_path: Path to the DOCX file
        
    Returns:
        Extracted text as string
    """
    docx_file = Path(docx_path)

    if not docx_file.exists():
        raise FileNotFoundError(f"DOCX file not found: {docx_path}")

    doc = Document(docx_file)

    print(f"📄 Document: {docx_file.name}")
    print(f"📊 Paragraphs: {len(doc.paragraphs)}")
    print(f"📊 Tables: {len(doc.tables)}")
    print("\n" + "="*80)
    print("EXTRACTING TEXT...")
    print("="*80 + "\n")

    text_content = []

    # Extract paragraphs
    for i, para in enumerate(doc.paragraphs, 1):
        if para.text.strip():
            text_content.append(para.text)
            if i % 50 == 0:
                print(f"Processed {i} paragraphs...")

    # Extract tables
    for table_num, table in enumerate(doc.tables, 1):
        text_content.append(f"\n\n[TABLE {table_num}]")
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip():
                text_content.append(row_text)

    full_text = "\n".join(text_content)
    return full_text

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python extract_docx.py <docx_file_path>")
        sys.exit(1)

    docx_path = sys.argv[1]

    try:
        extracted_text = extract_docx_text(docx_path)

        # Save to output file
        output_path = Path(docx_path).stem + "_extracted.txt"
        output_file = Path("/Users/reh3376/repos/plc-gbt/docs/FT") / output_path

        with open(output_file, 'w', encoding='utf-8') as f:
            f.write(extracted_text)

        print("\n✅ Text extracted successfully!")
        print(f"📄 Output saved to: {output_file}")
        print(f"📊 Total characters: {len(extracted_text):,}")

    except Exception as e:
        print(f"\n❌ Error: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)

